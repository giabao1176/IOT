"""Render the Word-exported PDF with the bundled PDFium runtime and audit navigation."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile
from lxml import etree
from docx import Document
import pypdfium2 as pdfium
from pypdf import PdfReader

BASE = Path(__file__).resolve().parent
NAME = 'Bao_Cao_Tieu_Luan_Cuoi_Khoa_C5_DangGiaHuy_HoanThien'

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--reviewed', action='store_true', help='Only after the agent visually inspected every generated page')
    args = parser.parse_args()
    target = BASE/'reports/delivery_document_qa.json'
    if args.reviewed:
        qa = json.loads(target.read_text(encoding='utf-8'))
        assert qa['docx_sha256'] == digest(BASE/(NAME+'.docx'))
        assert qa['pdf_sha256'] == digest(BASE/(NAME+'.pdf'))
        assert all((BASE/qa['preview_directory']/('page_%02d.png'%i)).is_file() for i in range(1,qa['pages']+1))
        qa['all_pages_visually_reviewed'] = True
        qa['review'] = 'All rendered pages inspected by the agent after the latest Word field update and PDF export.'
        target.write_text(json.dumps(qa,indent=2,ensure_ascii=False,allow_nan=False),encoding='utf-8')
        print(json.dumps(qa,ensure_ascii=False))
        return
    directory = BASE/'reports'/('qa_final_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
    directory.mkdir(parents=True,exist_ok=False)
    pdf = pdfium.PdfDocument(str(BASE/(NAME+'.pdf')))
    for i in range(len(pdf)):
        page = pdf[i]
        bitmap = page.render(scale=150/72)
        bitmap.to_pil().save(directory/('page_%02d.png'%(i+1)))
        bitmap.close()
        page.close()
    reader = PdfReader(BASE/(NAME+'.pdf'))
    texts = [p.extract_text() for p in reader.pages]
    headings = [(i+1, line) for i,t in enumerate(texts) for line in t.splitlines() if re.match(r'CHƯƠNG [1-6]',line)]
    doc = Document(BASE/(NAME+'.docx'))
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with ZipFile(BASE/(NAME+'.docx')) as z:
        root = etree.fromstring(z.read('word/document.xml'))
    bookmarks = set(root.xpath('//w:bookmarkStart/@w:name',namespaces=ns))
    anchors = root.xpath('//w:hyperlink/@w:anchor',namespaces=ns)
    fields = root.xpath('//w:fldSimple/@w:instr | //w:instrText/text()',namespaces=ns)
    refs = [m.group(1) for f in fields if (m:=re.search(r'REF\s+(\w+)',f))]
    fills = root.xpath('//w:shd/@w:fill',namespaces=ns)
    qa = dict(docx_sha256=digest(BASE/(NAME+'.docx')),pdf_sha256=digest(BASE/(NAME+'.pdf')),
              pages=len(reader.pages),chapter_heading_pages=headings,
              internal_anchors=len(anchors),equation_references=len(refs),
              broken_anchors=sorted((set(anchors)|set(refs))-bookmarks),
              native_equations=len(root.xpath('//m:oMath',namespaces=ns)),
              nonwhite_table_fills=[v for v in fills if v.lower() not in ['ffffff','auto']],
              leading_dash_paragraphs=[p.text for p in doc.paragraphs if p.text.lstrip().startswith('-')],
              sections=[dict(left_cm=s.left_margin.cm,right_cm=s.right_margin.cm,top_cm=s.top_margin.cm,bottom_cm=s.bottom_margin.cm) for s in doc.sections],
              all_pages_visually_reviewed=False,preview_directory=directory.relative_to(BASE).as_posix(),
              page_limit_user_override=True,
              page_limit_note='Proposal specifies 15–25 pages; user permits a longer report. No advisor approval of an exception is claimed.',
              rendering='Microsoft Word export followed by bundled PDFium at 150 dpi',
              training_origin_verified=False)
    assert not qa['broken_anchors'] and not qa['nonwhite_table_fills'] and not qa['leading_dash_paragraphs']
    assert all(abs(s['left_cm']-2.5)<.01 and abs(s['right_cm']-2)<.01 for s in qa['sections'])
    target.write_text(json.dumps(qa,indent=2,ensure_ascii=False,allow_nan=False),encoding='utf-8')
    print(json.dumps(qa,ensure_ascii=False,indent=2))

if __name__=='__main__':
    main()
