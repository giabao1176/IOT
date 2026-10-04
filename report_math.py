"""Native editable equations with Word sequence fields and internal references."""
import copy
import re
from xml.sax.saxutils import escape
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
REGISTRY = {}

def r(text):
    return '<m:r><m:t>' + escape(text) + '</m:t></m:r>'

def sub(base, index):
    return '<m:sSub><m:e>' + r(base) + '</m:e><m:sub>' + r(index) + '</m:sub></m:sSub>'

def fraction(top, bottom):
    return '<m:f><m:num>' + top + '</m:num><m:den>' + bottom + '</m:den></m:f>'

def custom_math(kind):
    if kind == 'DCT':
        summed = sub('x','n') + r('cos') + '<m:d><m:e>' + fraction(r('π(2n + 1)k'), r('2N')) + '</m:e></m:d>'
        body = sub('X','k') + r(' = ') + sub('α','k')
        body += '<m:nary><m:naryPr><m:chr m:val="∑"/><m:limLoc m:val="undOvr"/></m:naryPr><m:sub>' + r('n = 0') + '</m:sub><m:sup>' + r('N − 1') + '</m:sup><m:e>' + summed + '</m:e></m:nary>'
    elif kind == 'STE':
        body = fraction(r('∂') + '<m:acc><m:accPr><m:chr m:val="̃"/></m:accPr><m:e>' + r('z') + '</m:e></m:acc>', r('∂z')) + r(' ≈ 1')
    elif kind == 'RR':
        body = sub('R','a') + r(' = ') + fraction(r('60(') + sub('K','a') + r(' − 1)'), sub('t','a,K') + r(' − ') + sub('t','a,1'))
    else:
        raise ValueError(kind)
    return parse_xml('<m:oMath xmlns:m="' + NS + '">' + body + '</m:oMath>')

def add_equation(doc, math, key):
    number = len(REGISTRY) + 1
    bookmark = 'EQ_' + str(number)
    REGISTRY[key] = (number, bookmark)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p._p.append(copy.deepcopy(math))
    p.add_run('   (')
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(10000 + number))
    start.set(qn('w:name'), bookmark)
    p._p.append(start)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'SEQ Equation \\* ARABIC')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = str(number)
    run.append(text)
    field.append(run)
    p._p.append(field)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(10000 + number))
    p._p.append(end)
    p.add_run(')')
    return p

def clone_equation(doc, source, paragraph_idx):
    paragraph = source.paragraphs[paragraph_idx]
    maths = paragraph._p.findall('.//' + qn('m:oMath'))
    if not maths:
        raise ValueError('Source paragraph lacks native equation: ' + str(paragraph_idx))
    result = None
    for original in maths:
        math = copy.deepcopy(original)
        last = math[-1]
        label = ''.join(last.itertext()) if last.tag == qn('m:d') else ''
        # The proposal stores its equation label as a trailing math delimiter.
        values = last.findall('.//' + qn('m:t'))
        key = ''.join(t.text or '' for t in values).strip()
        if not key.isdigit():
            raise ValueError('Cannot identify source formula label')
        math.remove(last)
        while len(math) and math[-1].tag == qn('m:r') and ''.join(t.text or '' for t in math[-1].findall('.//' + qn('m:t'))).strip() in ['', '.']:
            math.remove(math[-1])
        result = add_equation(doc, math, key)
    return result

def update_formula_references(doc):
    pattern = re.compile(r'công thức \(([0-9., ]+)\)', re.IGNORECASE)
    paragraphs = list(doc.paragraphs) + [p for t in doc.tables for row in t.rows for c in row.cells for p in c.paragraphs]
    for p in paragraphs:
        content = p.text
        matches = list(pattern.finditer(content))
        if not matches:
            continue
        p.clear()
        cursor = 0
        for match in matches:
            p.add_run(content[cursor:match.start()] + 'công thức (')
            keys = [x.strip() for x in match.group(1).split(',')]
            for i, key in enumerate(keys):
                if key not in REGISTRY:
                    raise ValueError('Unresolved equation reference: ' + key)
                if i:
                    p.add_run(', ')
                number, bookmark = REGISTRY[key]
                field = OxmlElement('w:fldSimple')
                field.set(qn('w:instr'), 'REF ' + bookmark + ' \\h')
                run = OxmlElement('w:r')
                text = OxmlElement('w:t')
                text.text = str(number)
                run.append(text)
                field.append(run)
                p._p.append(field)
            p.add_run(')')
            cursor = match.end()
        p.add_run(content[cursor:])
