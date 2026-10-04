"""Read-only DOCX content and formatting inventory for project review."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
from lxml import etree
from docx import Document
import sys

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
for name in sys.argv[1:]:
    path = Path(name)
    doc = Document(path)
    print("\nDOCUMENT", path.name)
    for s in doc.sections:
        print("SECTION_XML", etree.tostring(s._sectPr, encoding="unicode"))
    print("STYLES", Counter(p.style.name for p in doc.paragraphs))
    print("RUN_FONTS", Counter((r.font.name, r.font.size.pt if r.font.size else None)
                               for p in doc.paragraphs for r in p.runs if r.text.strip()))
    print("SPACING", Counter(str(p.paragraph_format.line_spacing) for p in doc.paragraphs))
    with ZipFile(path) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        print("FIELDS", root.xpath('//w:instrText/text()', namespaces=NS))
        print("SHADING", Counter(root.xpath('//w:shd/@w:fill', namespaces=NS)))
        print("BORDER_COLORS", Counter(root.xpath('//w:tblBorders/*/@w:color', namespaces=NS)))
        for p in root.xpath('//w:body//w:p', namespaces=NS):
            text = ''.join(p.xpath('.//w:t/text() | .//*[local-name()="t" and namespace-uri()="http://schemas.openxmlformats.org/officeDocument/2006/math"]/text()', namespaces=NS))
            if text.strip():
                print(text)
